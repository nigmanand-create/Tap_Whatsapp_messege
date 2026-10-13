from __future__ import annotations

from datetime import time

import frappe
from frappe.utils import getdate, get_time

from tap_buddy.wgas.excel_utils import combine_schedule_local, map_headers, split_groups
from tap_buddy.wgas.resolvers import resolve_glific_flow, resolve_whatsapp_group


def import_excel_into_plan(plan_name: str) -> str:
	from openpyxl import load_workbook

	plan = frappe.get_doc("WGAS Plan", plan_name)
	if not plan.excel_file:
		frappe.throw("Excel file is required.")

	file_name = frappe.db.get_value("File", {"file_url": plan.excel_file}, "name")
	if not file_name:
		frappe.throw("Could not resolve uploaded Excel file.")
	file_doc = frappe.get_doc("File", file_name)
	file_path = file_doc.get_full_path()

	wb = load_workbook(filename=file_path, read_only=True, data_only=True)
	ws = wb.active
	rows = list(ws.iter_rows(values_only=True))
	if not rows:
		frappe.throw("Excel sheet is empty.")

	headers = map_headers(list(rows[0]))
	missing = [k for k in ("date", "time", "groups") if k not in headers]
	if missing:
		frappe.throw(
			f"Missing required columns: {', '.join(missing)}. "
			"Expected headers like date, time, whatsapp group(s), message, flow."
		)

	plan.items = []
	item_count = 0
	skipped = 0
	tz_name = plan.timezone or frappe.db.get_single_value("WGAS Settings", "default_timezone") or "Asia/Kolkata"

	for row_idx, row in enumerate(rows[1:], start=2):
		if not row or all(cell is None or str(cell).strip() == "" for cell in row):
			continue

		row_date = row[headers["date"]]
		row_time = row[headers["time"]]
		groups_raw = row[headers["groups"]]
		message = row[headers.get("message", -1)] if "message" in headers else None
		flow_ref = row[headers.get("flow", -1)] if "flow" in headers else None

		message_text = str(message).strip() if message is not None else ""
		flow_text = str(flow_ref).strip() if flow_ref is not None else ""

		if not message_text and not flow_text:
			skipped += 1
			continue

		scheduled_at = combine_schedule_local(row_date, row_time, tz_name)
		groups = split_groups(groups_raw)
		if not groups:
			skipped += 1
			continue

		for group_label in groups:
			wa_group = resolve_whatsapp_group(group_label)
			glific_flow = resolve_glific_flow(flow_text) if flow_text else None
			mode = "Flow" if flow_text else "Text Message"
			status = "Pending"
			err = None
			if not wa_group:
				status = "Skipped"
				err = f"Group not found: {group_label}"
			elif flow_text and not glific_flow:
				status = "Skipped"
				err = f"Flow not found: {flow_text}"

			plan.append(
				"items",
				{
					"source_row": row_idx,
					"schedule_date": getdate(row_date),
					"schedule_time": get_time(row_time) if row_time else time(9, 0),
					"scheduled_at": scheduled_at,
					"group_label": group_label,
					"whatsapp_group": wa_group,
					"message": message_text,
					"flow_ref": flow_text,
					"glific_flow": glific_flow,
					"dispatch_mode": mode,
					"status": status,
					"error_message": err,
				},
			)
			item_count += 1

	plan.import_summary = f"Imported {item_count} schedule item(s); skipped {skipped} row(s)."
	plan.status = "Imported"
	plan.save(ignore_permissions=True)
	frappe.db.commit()

	from tap_buddy.tap_buddy.doctype.wgas_plan.wgas_plan import update_plan_stats

	update_plan_stats(plan.name)
	return plan.import_summary
